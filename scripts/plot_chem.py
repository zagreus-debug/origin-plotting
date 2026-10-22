#!/usr/bin/env python3
"""Origin-style chemistry plots with optional automatic model selection.

No Origin, NumPy, SciPy, Matplotlib, or pandas is required. Only Pillow is
needed; openpyxl is optional for .xlsx input.
"""

from __future__ import annotations

import argparse
import csv
import math
import sys
from dataclasses import dataclass
from pathlib import Path

try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    print("ERROR: Pillow is required. Install it with: pip install pillow", file=sys.stderr)
    raise SystemExit(1)


ALIASES = {
    "": "auto", "auto": "auto", "automatic": "auto", "自动": "auto",
    "none": "none", "no": "none", "off": "none", "不拟合": "none",
    "linear": "linear", "line": "linear", "线性": "linear",
    "quad": "quad", "quadratic": "quad", "poly2": "quad", "二次": "quad",
    "exp": "exp", "exponential": "exp", "指数": "exp",
    "power": "power", "幂": "power", "幂函数": "power",
    "log": "log", "logarithmic": "log", "对数": "log",
    "langmuir": "langmuir", "michaelis": "langmuir", "饱和": "langmuir",
}

MODEL_NAMES = {
    "linear": "linear", "quad": "quadratic", "exp": "exponential",
    "power": "power", "log": "logarithmic", "langmuir": "Langmuir",
}


@dataclass
class Fit:
    key: str
    params: list[tuple[str, float]]
    equation: str
    r2: float
    aicc: float
    predict: object


def solve_linear(a: list[list[float]], b: list[float]) -> list[float]:
    n = len(b)
    m = [[float(a[i][j]) for j in range(n)] + [float(b[i])] for i in range(n)]
    for i in range(n):
        pivot = max(range(i, n), key=lambda r: abs(m[r][i]))
        if abs(m[pivot][i]) < 1e-14:
            raise ValueError("Singular fit matrix.")
        if pivot != i:
            m[i], m[pivot] = m[pivot], m[i]
        pv = m[i][i]
        for j in range(i, n + 1):
            m[i][j] /= pv
        for r in range(n):
            if r == i:
                continue
            factor = m[r][i]
            if factor:
                for j in range(i, n + 1):
                    m[r][j] -= factor * m[i][j]
    return [m[i][n] for i in range(n)]


def polyfit_normal(x: list[float], y: list[float], degree: int) -> list[float]:
    scale = max(abs(v) for v in x) or 1.0
    z = [v / scale for v in x]
    size = degree + 1
    a = [[sum(zi ** (i + j) for zi in z) for j in range(size)] for i in range(size)]
    b = [sum((zi ** i) * yi for zi, yi in zip(z, y)) for i in range(size)]
    coef_z = solve_linear(a, b)
    coef_x = [coef_z[i] / (scale ** i) for i in range(size)]
    return coef_x[::-1]


def aicc(sse: float, n: int, k: int) -> float:
    if sse <= 0:
        sse = 1e-300
    if n <= k + 1:
        return math.inf
    return n * math.log(sse / n) + 2 * k + (2 * k * (k + 1)) / (n - k - 1)


def make_fit(key: str, x: list[float], y: list[float]) -> Fit | None:
    n = len(x)
    if key in ("linear", "quad"):
        degree = 1 if key == "linear" else 2
        if n < degree + 1:
            return None
        coef = polyfit_normal(x, y, degree)
        if key == "linear":
            a, b = coef
            predict = lambda v, a=a, b=b: a * v + b
            equation = f"y = {a:.4g} x + {b:.4g}"
            params = [("a", a), ("b", b)]
        else:
            a, b, c = coef
            predict = lambda v, a=a, b=b, c=c: a * v * v + b * v + c
            equation = f"y = {a:.4g} x^2 + {b:.4g} x + {c:.4g}"
            params = [("a", a), ("b", b), ("c", c)]
    elif key == "exp":
        if n < 3:
            return None
        if all(v > 0 for v in y):
            sign = 1.0
        elif all(v < 0 for v in y):
            sign = -1.0
        else:
            return None
        b, ln_a = polyfit_normal(x, [math.log(sign * v) for v in y], 1)
        a = sign * math.exp(ln_a)
        predict = lambda v, a=a, b=b: a * math.exp(b * v)
        equation = f"y = {a:.4g} exp({b:.4g} x)"
        params = [("a", a), ("b", b)]
    elif key == "power":
        if n < 3 or any(v <= 0 for v in x) or any(v <= 0 for v in y):
            return None
        b, ln_a = polyfit_normal([math.log(v) for v in x], [math.log(v) for v in y], 1)
        a = math.exp(ln_a)
        predict = lambda v, a=a, b=b: a * (v ** b)
        equation = f"y = {a:.4g} x^{b:.4g}"
        params = [("a", a), ("b", b)]
    elif key == "log":
        if n < 3 or any(v <= 0 for v in x):
            return None
        b, a = polyfit_normal([math.log(v) for v in x], y, 1)
        predict = lambda v, a=a, b=b: a + b * math.log(v)
        equation = f"y = {a:.4g} + {b:.4g} ln(x)"
        params = [("a", a), ("b", b)]
    elif key == "langmuir":
        if n < 3 or any(v <= 0 for v in x) or any(v <= 0 for v in y):
            return None
        m, c = polyfit_normal([1 / v for v in x], [1 / v for v in y], 1)
        if c == 0:
            return None
        a = 1 / c
        b = m / c
        predict = lambda v, a=a, b=b: a * v / (b + v)
        equation = f"y = {a:.4g} x / ({b:.4g} + x)"
        params = [("a", a), ("b", b)]
    else:
        return None

    sse = sum((yi - predict(xi)) ** 2 for xi, yi in zip(x, y))
    sst = sum((yi - sum(y) / n) ** 2 for yi in y)
    r2 = 1 - sse / sst if sst > 0 else 1.0
    return Fit(key, params, equation, r2, aicc(sse, n, len(params)), predict)


def select_fit(requested: str, x: list[float], y: list[float]) -> Fit | None:
    key = ALIASES.get(requested.strip().lower())
    if key is None:
        raise ValueError(f"Unknown fit model '{requested}'.")
    if key == "none":
        return None
    if key != "auto":
        fit = make_fit(key, x, y)
        if fit is None:
            raise ValueError(f"Model '{key}' cannot fit this data (check sign/zero constraints).")
        return fit
    candidates = [f for f in (make_fit(k, x, y) for k in MODEL_NAMES) if f is not None]
    if not candidates:
        raise ValueError("No candidate model fits this data.")
    best = min(candidates, key=lambda f: f.aicc)
    linear = next((f for f in candidates if f.key == "linear"), None)
    if linear and linear.aicc <= best.aicc + 2:
        best = linear
    return best


def read_table(path: str, sheet: str | None) -> tuple[list[str], list[dict]]:
    suffix = Path(path).suffix.lower()
    if suffix in (".csv", ".tsv", ".txt"):
        delimiter = "\t" if suffix == ".tsv" else ","
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle, delimiter=delimiter)
            headers = [str(h).strip() for h in (reader.fieldnames or [])]
            rows = [{str(k).strip(): v for k, v in row.items()} for row in reader]
        return headers, rows
    if suffix in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("openpyxl is required for xlsx input: pip install openpyxl") from exc
        wb = load_workbook(path, data_only=True, read_only=True)
        ws = wb[sheet] if sheet else wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        headers = [str(v).strip() if v is not None else "" for v in next(rows_iter)]
        rows = [dict(zip(headers, values)) for values in rows_iter]
        return headers, rows
    raise ValueError("Supported inputs: .csv, .tsv, .txt, .xlsx, .xlsm")


def parse_float(value) -> float | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def nice_number(value: float, round_up: bool = False) -> float:
    if value <= 0:
        return 1.0
    exp = math.floor(math.log10(value))
    frac = value / (10 ** exp)
    if round_up:
        nf = 1 if frac < 1.5 else 2 if frac < 3 else 5 if frac < 7 else 10
    else:
        nf = 1 if frac <= 1 else 2 if frac <= 2 else 5 if frac <= 5 else 10
    return nf * (10 ** exp)


def axis_ticks(lo: float, hi: float, target: int = 6) -> tuple[list[float], float]:
    step = nice_number((hi - lo) / target, True)
    start = math.ceil(lo / step - 1e-10) * step
    ticks = []
    v = start
    while v <= hi + step * 1e-6:
        ticks.append(v if abs(v) > step * 1e-9 else 0.0)
        v += step
    if len(ticks) < 2:
        ticks = [lo, hi]
        step = hi - lo
    return ticks, step


def format_tick(value: float, step: float) -> str:
    if abs(value) < 1e-12:
        value = 0.0
    if abs(value) >= 1e4 or (abs(value) > 0 and abs(value) < 1e-3):
        return f"{value:.2e}"
    decimals = max(0, -int(math.floor(math.log10(step))))
    return f"{value:.{decimals}f}"


FONT_CANDIDATES = {
    "serif": [
        "C:/Windows/Fonts/times.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf",
        "/System/Library/Fonts/Times.ttc",
    ],
    "serif_bold": [
        "C:/Windows/Fonts/timesbd.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf",
        "/System/Library/Fonts/Times.ttc",
    ],
    "cjk": [
        "C:/Windows/Fonts/simhei.ttf",
        "C:/Windows/Fonts/msyh.ttc",
        "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc",
        "/usr/share/fonts/truetype/wqy/wqy-zenhei.ttc",
        "/System/Library/Fonts/PingFang.ttc",
    ],
}
FONT_CACHE: dict[tuple[str, int], ImageFont.FreeTypeFont] = {}


def has_cjk(text: str) -> bool:
    return any("\u4e00" <= ch <= "\u9fff" or "\u3000" <= ch <= "\u303f" for ch in text)


def get_font(text: str, size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    kind = "cjk" if has_cjk(text) else ("serif_bold" if bold else "serif")
    for path in FONT_CANDIDATES[kind]:
        if Path(path).exists():
            key = (path, size)
            if key not in FONT_CACHE:
                FONT_CACHE[key] = ImageFont.truetype(path, size)
            return FONT_CACHE[key]
    return ImageFont.load_default()


def text_size(draw: ImageDraw.ImageDraw, text: str, font) -> tuple[int, int]:
    box = draw.textbbox((0, 0), text, font=font)
    return box[2] - box[0], box[3] - box[1]


def make_plot(headers: list[str], rows: list[dict], args: argparse.Namespace) -> tuple[Path, Path, Path]:
    for col in (args.x, args.y, args.yerr):
        if col and col not in headers:
            raise ValueError(f"Column '{col}' not found. Available: {', '.join(headers)}")

    points = []
    dropped = 0
    for row in rows:
        x = parse_float(row.get(args.x))
        y = parse_float(row.get(args.y))
        e = parse_float(row.get(args.yerr)) if args.yerr else None
        if x is None or y is None or (args.yerr and (e is None or e < 0)):
            dropped += 1
            continue
        points.append((x, y, e))
    if len(points) < 2:
        raise ValueError("Need at least 2 valid X/Y rows.")
    points.sort(key=lambda p: p[0])
    x = [p[0] for p in points]
    y = [p[1] for p in points]
    if max(x) == min(x):
        raise ValueError("All X values are equal; cannot fit or plot a line.")
    fit = select_fit(args.fit, x, y)

    dpi = args.dpi
    unit = dpi / 300.0
    sup = 2
    W = int(round(8 * dpi)) * sup
    H = int(round(6 * dpi)) * sup

    def U(v: float) -> int:
        return int(round(v * unit * sup))

    img = Image.new("RGB", (W, H), "white")
    draw = ImageDraw.Draw(img)
    L, T, R, B = U(300), U(200), W - U(110), H - U(250)

    ylo = min(y)
    yhi = max(y)
    if args.yerr:
        ylo = min(ylo, min(v - e for _, v, e in points))
        yhi = max(yhi, max(v + e for _, v, e in points))
    xpad = (max(x) - min(x)) * 0.06
    ypad = (yhi - ylo) * 0.08 or 1.0
    xmin, xmax = min(x) - xpad, max(x) + xpad
    ymin, ymax = ylo - ypad, yhi + ypad

    def X(v: float) -> float:
        return L + (v - xmin) / (xmax - xmin) * (R - L)

    def Y(v: float) -> float:
        return B - (v - ymin) / (ymax - ymin) * (B - T)

    if args.title:
        font = get_font(args.title, U(46))
        tw, th = text_size(draw, args.title, font)
        draw.text(((W - tw) / 2, U(46)), args.title, font=font, fill=(17, 17, 17))

    draw.rectangle([L, T, R, B], outline=(17, 17, 17), width=max(1, U(2)))
    xticks, xstep = axis_ticks(xmin, xmax)
    yticks, ystep = axis_ticks(ymin, ymax, 7)
    tick_font = get_font("0", U(28))
    major = U(12)
    minor = U(7)

    for v in xticks:
        px = X(v)
        draw.line([px, B, px, B - major], fill=(17, 17, 17), width=max(1, U(2)))
        draw.line([px, T, px, T + major], fill=(17, 17, 17), width=max(1, U(2)))
        label = format_tick(v, xstep)
        tw, th = text_size(draw, label, tick_font)
        draw.text((px - tw / 2, B + U(18)), label, font=tick_font, fill=(17, 17, 17))
    for v in yticks:
        py = Y(v)
        draw.line([L, py, L + major, py], fill=(17, 17, 17), width=max(1, U(2)))
        draw.line([R, py, R - major, py], fill=(17, 17, 17), width=max(1, U(2)))
        label = format_tick(v, ystep)
        tw, th = text_size(draw, label, tick_font)
        draw.text((L - U(18) - tw, py - th / 2 - U(3)), label, font=tick_font, fill=(17, 17, 17))

    xminor = xstep / 5
    v = math.ceil(xmin / xminor) * xminor
    while v <= xmax:
        px = X(v)
        draw.line([px, B, px, B - minor], fill=(17, 17, 17), width=max(1, U(1.5)))
        draw.line([px, T, px, T + minor], fill=(17, 17, 17), width=max(1, U(1.5)))
        v += xminor
    yminor = ystep / 5
    v = math.ceil(ymin / yminor) * yminor
    while v <= ymax:
        py = Y(v)
        draw.line([L, py, L + minor, py], fill=(17, 17, 17), width=max(1, U(1.5)))
        draw.line([R, py, R - minor, py], fill=(17, 17, 17), width=max(1, U(1.5)))
        v += yminor

    xlabel = args.xlabel or args.x
    ylabel = args.ylabel or args.y
    xfont = get_font(xlabel, U(36))
    tw, th = text_size(draw, xlabel, xfont)
    draw.text(((L + R - tw) / 2, B + U(88)), xlabel, font=xfont, fill=(17, 17, 17))
    yfont = get_font(ylabel, U(36))
    tw, th = text_size(draw, ylabel, yfont)
    tmp = Image.new("RGBA", (tw + U(20), th + U(20)), (255, 255, 255, 0))
    ImageDraw.Draw(tmp).text((U(10), U(10)), ylabel, font=yfont, fill=(17, 17, 17))
    tmp = tmp.rotate(90, expand=True)
    img.paste(tmp, (U(58), int((T + B) / 2 - tmp.height / 2)), tmp)

    layer = Image.new("RGBA", (W, H), (255, 255, 255, 0))
    ld = ImageDraw.Draw(layer)
    red = (214, 39, 40, 255)
    dark_red = (127, 29, 29, 255)
    blue = (31, 78, 156, 255)
    r = U(7)

    if args.yerr:
        for xv, yv, ev in points:
            px, py = X(xv), Y(yv)
            y0, y1 = Y(yv - ev), Y(yv + ev)
            ld.line([px, y0, px, y1], fill=(51, 51, 51, 255), width=max(1, U(1.5)))
            ld.line([px - U(8), y0, px + U(8), y0], fill=(51, 51, 51, 255), width=max(1, U(1.5)))
            ld.line([px - U(8), y1, px + U(8), y1], fill=(51, 51, 51, 255), width=max(1, U(1.5)))

    if args.style in ("line", "line+symbol"):
        coords = [(X(xv), Y(yv)) for xv, yv, _ in points]
        ld.line(coords, fill=red, width=max(1, U(2)))
    if args.style in ("scatter", "line+symbol"):
        for xv, yv, _ in points:
            px, py = X(xv), Y(yv)
            ld.ellipse([px - r, py - r, px + r, py + r], fill=red, outline=dark_red,
                       width=max(1, U(1.5)))

    if fit is not None:
        samples = [(L + (R - L) * i / 399.0, 0.0) for i in range(400)]
        coords = []
        for px, _ in samples:
            xv = xmin + (px - L) / (R - L) * (xmax - xmin)
            coords.append((px, Y(fit.predict(xv))))
        ld.line(coords, fill=blue, width=max(1, U(2.4)))

    mask = Image.new("L", (W, H), 0)
    ImageDraw.Draw(mask).rectangle([L, T, R, B], fill=255)
    clipped = Image.new("RGBA", (W, H), (0, 0, 0, 0))
    clipped.paste(layer, (0, 0), mask)
    img = Image.alpha_composite(img.convert("RGBA"), clipped).convert("RGB")
    draw = ImageDraw.Draw(img)

    entries = [(args.data_label, "data")]
    if fit is not None:
        entries.append((args.fit_label, "fit"))
    legend_font = get_font("Data", U(26))
    widths = [text_size(draw, label, legend_font)[0] for label, _ in entries]
    lw = max(widths) + U(110)
    lh = U(34) * len(entries) + U(30)
    lx0, ly0 = R - lw - U(20), T + U(20)
    draw.rectangle([lx0, ly0, lx0 + lw, ly0 + lh], fill="white", outline=(136, 136, 136), width=max(1, U(1.5)))
    for i, (label, kind) in enumerate(entries):
        cy = ly0 + U(26) + i * U(34)
        if kind == "data":
            draw.ellipse([lx0 + U(24) - r, cy - r, lx0 + U(24) + r, cy + r], fill=red, outline=dark_red,
                         width=max(1, U(1.5)))
        else:
            draw.line([lx0 + U(12), cy, lx0 + U(36), cy], fill=blue, width=max(1, U(2.4)))
        f = get_font(label, U(26))
        draw.text((lx0 + U(56), cy - U(13)), label, font=f, fill=(17, 17, 17))

    if fit is not None:
        note_font = get_font(fit.equation, U(26))
        lines = [fit.equation, f"R2 = {fit.r2:.6f}"]
        ny = B - U(100)
        for i, line in enumerate(lines):
            f = get_font(line, U(26))
            draw.text((L + U(45), ny + i * U(42)), line, font=f, fill=(17, 17, 17))

    outdir = Path(args.outdir)
    outdir.mkdir(parents=True, exist_ok=True)
    png = outdir / f"{args.name}.png"
    pdf = outdir / f"{args.name}.pdf"
    report = outdir / f"{args.name}_fit.txt"
    final = img.resize((W // sup, H // sup), Image.LANCZOS)
    final.save(png, dpi=(dpi, dpi))
    final.convert("RGB").save(pdf, "PDF", resolution=dpi)

    lines = [
        f"model: {fit.key if fit else 'none'}",
        f"model_name: {MODEL_NAMES.get(fit.key, 'none') if fit else 'none'}",
        f"n: {len(points)}",
        f"dropped_rows: {dropped}",
    ]
    if fit:
        lines += [f"equation: {fit.equation}", f"r2: {fit.r2:.8f}"]
        lines += [f"{k}: {v:.10g}" for k, v in fit.params]
    report.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return png, pdf, report


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--input", required=True)
    p.add_argument("--x", required=True)
    p.add_argument("--y", required=True)
    p.add_argument("--yerr")
    p.add_argument("--sheet")
    p.add_argument("--fit", default="auto", help="auto, none, linear, quad, exp, power, log, langmuir")
    p.add_argument("--title", default="")
    p.add_argument("--xlabel", default="")
    p.add_argument("--ylabel", default="")
    p.add_argument("--style", choices=["scatter", "line", "line+symbol"], default="scatter")
    p.add_argument("--data-label", default="Data")
    p.add_argument("--fit-label", default="Fit")
    p.add_argument("--outdir", default="chem_output")
    p.add_argument("--name", default="chem_plot")
    p.add_argument("--dpi", type=int, default=300)
    return p.parse_args()


def main() -> int:
    args = parse_args()
    headers, rows = read_table(args.input, args.sheet)
    png, pdf, report = make_plot(headers, rows, args)
    print(f"Saved: {png}")
    print(f"Saved: {pdf}")
    print(f"Saved: {report}")
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        raise SystemExit(1)
